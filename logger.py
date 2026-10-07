import os
from datetime import datetime
from openpyxl import Workbook, load_workbook

FILE = "bets.xlsx"

def init_excel():
    if not os.path.exists(FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Bets"
        ws.append(["Datum loggad", "Matchdatum", "Sport", "Match", "Lag", "Marknad", "Bookmaker", "Odds", "Rättvisa odds", "Edge %", "Units", "Insats Flat (kr)", "Insats Kelly (kr)", "Resultat", "V/F Flat", "V/F Kelly"])
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 35
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 10  # Marknad
        ws.column_dimensions['G'].width = 15  # Bookmaker
        ws.column_dimensions['H'].width = 8   # Odds
        ws.column_dimensions['I'].width = 14  # Rättvisa odds
        ws.column_dimensions['J'].width = 8   # Edge %
        ws.column_dimensions['K'].width = 8   # Units
        ws.column_dimensions['L'].width = 14  # Insats Flat
        ws.column_dimensions['M'].width = 14  # Insats Kelly
        ws.column_dimensions['N'].width = 10  # Resultat
        ws.column_dimensions['O'].width = 12  # V/F Flat
        ws.column_dimensions['P'].width = 12  # V/F Kelly
        ws.column_dimensions['Q'].width = 3   # mellanrum
        wb.save(FILE)

def log_bet(match, team, bookmaker, offered_odds, fair, edge, units, sport, market_type):
    match_date = match['commence_time'][:10]
    logged_date = datetime.now().strftime("%Y-%m-%d")
    kelly_stake = round(units * 10, 2)

    wb = load_workbook(FILE)
    ws = wb.active

    last_row = ws.max_row
    if last_row > 1:
        last_date = ws.cell(row=last_row, column=1).value
        if last_date and str(last_date) != logged_date:
            ws.append([""])

    row = ws.max_row + 1

    ws.append([
        logged_date,                                      # A
        match_date,                                       # B
        sport,                                            # C
        f"{match['home_team']} vs {match['away_team']}",  # D
        team,                                             # E: Lag
        market_type,                                      # F: Marknad (h2h/totals)
        bookmaker['title'],                               # G: Bookmaker
        offered_odds,                                     # H: Odds
        round(fair, 2),                                   # I: Rättvisa odds
        round(edge * 100, 1),                             # J: Edge %
        units,                                            # K: Units
        20,                                               # L: Insats Flat
        kelly_stake,                                      # M: Insats Kelly
        "",                                               # N: Resultat
        f"=IF(UPPER(N{row})=\"W\",(H{row}-1)*L{row},IF(UPPER(N{row})=\"L\",-L{row},\"\"))",  # O: V/F Flat
        f"=IF(UPPER(N{row})=\"W\",(H{row}-1)*M{row},IF(UPPER(N{row})=\"L\",-M{row},\"\"))",  # P: V/F Kelly
    ])

    wb.save(FILE)