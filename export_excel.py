import os
import sqlite3
import sys
from contextlib import closing
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

DB_FILE = "bets.db"
OUTPUT_FILE = "bets_export.xlsx"
PROTECTED_FILE = "bets.xlsx"      # gamla loggen, får aldrig skrivas över

HEADERS = ["Datum loggad", "Matchdatum", "Sport", "Match", "Lag", "Marknad", "Bookmaker", "Odds",
           "Rättvisa odds", "Edge %", "Units", "Insats Flat (kr)", "Insats Kelly (kr)", "Resultat",
           "V/F Flat", "V/F Kelly"]

WIDTHS = {"A": 14, "B": 12, "C": 15, "D": 35, "E": 20, "F": 10, "G": 15, "H": 8, "I": 14,
          "J": 8, "K": 8, "L": 14, "M": 14, "N": 10, "O": 12, "P": 12, "Q": 3}

BOOKMAKER_NAMES = {
    "unibet_se": "Unibet (SE)",
    "pinnacle": "Pinnacle",
}


class ExportError(Exception):
    pass


def bookmaker_name(key):
    return BOOKMAKER_NAMES.get(key, key)


def outcome_label(name, point):
    """Samma format som gamla loggen: "Over 2.5" för totals, bara namnet för h2h."""
    return name if point is None else f"{name} {point:g}"


def local_date(logged_date):
    """logged_date sparas i UTC; datumet visas och grupperas i lokal tid."""
    return datetime.fromisoformat(logged_date).astimezone().strftime("%Y-%m-%d")


def vf_formula(row, stake_col):
    """V/F för en rad: W = (odds-1)*insats, L = -insats, P = 0 (insatsen tillbaka)."""
    result = f"UPPER(N{row})"
    return (f'=IF({result}="W",(H{row}-1)*{stake_col}{row},'
            f'IF({result}="L",-{stake_col}{row},'
            f'IF({result}="P",0,"")))')


def load_bets(db_file):
    if not os.path.exists(db_file):
        raise ExportError(f"Hittar inte {db_file} – har något bet loggats än?")
    # Skrivskyddat, så exporten aldrig kan ändra databasen
    uri = Path(db_file).resolve().as_uri() + "?mode=ro"
    with closing(sqlite3.connect(uri, uri=True)) as conn:
        conn.row_factory = sqlite3.Row
        return conn.execute("SELECT * FROM bets ORDER BY id").fetchall()


def build_workbook(bets):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bets"
    ws.append(HEADERS)
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width

    row = 2
    previous_date = None
    for bet in bets:
        date = local_date(bet["logged_date"])
        if previous_date is not None and date != previous_date:
            row += 1                                      # tomrad mellan loggdatum
        previous_date = date

        values = [
            date,                                         # A: Datum loggad
            bet["match_date"],                            # B: Matchdatum
            bet["sport_key"],                             # C: Sport
            bet["match_name"],                            # D: Match
            outcome_label(bet["outcome_name"], bet["point"]),  # E: Lag
            bet["market_key"],                            # F: Marknad
            bookmaker_name(bet["bookmaker"]),             # G: Bookmaker
            bet["offered_odds"],                          # H: Odds
            round(bet["fair_at_flag"], 2),                # I: Rättvisa odds
            round(bet["edge"] * 100, 1),                  # J: Edge %
            bet["units"],                                 # K: Units
            bet["stake_flat"],                            # L: Insats Flat
            bet["stake_kelly"],                           # M: Insats Kelly
            bet["result"],                                # N: Resultat
            vf_formula(row, "L"),                         # O: V/F Flat
            vf_formula(row, "M"),                         # P: V/F Kelly
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    last = max(row - 1, 2)                                # sista dataraden
    n = f"N2:N{last}"
    top = row + 1                                         # en tomrad före summeringen
    ws.cell(row=top, column=4, value="Summering").font = Font(bold=True)

    summary = [
        ("Antal bets", f"=COUNT(H2:H{last})", None),
        ("Antal avgjorda", f'=COUNTIF({n},"W")+COUNTIF({n},"L")+COUNTIF({n},"P")', None),
        ("Total V/F Flat (kr)", f"=SUM(O2:O{last})", "0.00"),
        ("Total V/F Kelly (kr)", f"=SUM(P2:P{last})", "0.00"),
        # ROI = vinst / insats på W+L; push räknas inte som omsättning
        ("ROI Flat", f'=IFERROR(E{top + 3}/(SUMIF({n},"W",L2:L{last})+SUMIF({n},"L",L2:L{last})),"")', "0.0%"),
        ("ROI Kelly", f'=IFERROR(E{top + 4}/(SUMIF({n},"W",M2:M{last})+SUMIF({n},"L",M2:M{last})),"")', "0.0%"),
    ]
    for i, (label, formula, number_format) in enumerate(summary, start=1):
        ws.cell(row=top + i, column=4, value=label)
        cell = ws.cell(row=top + i, column=5, value=formula)
        if number_format:
            cell.number_format = number_format

    return wb


def check_target(out_file):
    target = Path(out_file).resolve()
    if target == Path(PROTECTED_FILE).resolve() or target.name == PROTECTED_FILE:
        raise ExportError(f"Vägrar skriva över {PROTECTED_FILE}. Välj ett annat filnamn.")
    # Excel på Mac låser inte filen men skapar en lås-fil bredvid den
    for lock_name in (f"~${target.name}", f"~${target.name[2:]}"):
        if (target.parent / lock_name).exists():
            raise ExportError(f"{target.name} är öppen i Excel. Stäng den och kör igen.")


def export(db_file=DB_FILE, out_file=OUTPUT_FILE):
    """Skriver bets.db till out_file. Returnerar antal exporterade bets."""
    check_target(out_file)
    bets = load_bets(db_file)
    wb = build_workbook(bets)
    try:
        wb.save(out_file)
    except PermissionError:
        # Windows låser filen när den är öppen i Excel
        raise ExportError(f"{Path(out_file).name} är öppen i Excel. Stäng den och kör igen.")
    return len(bets)


if __name__ == "__main__":
    try:
        count = export()
    except ExportError as e:
        print(f"Fel: {e}", file=sys.stderr)
        sys.exit(1)
    print(f"Exporterade {count} bets till {OUTPUT_FILE}")
