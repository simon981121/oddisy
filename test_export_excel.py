import os
import sqlite3
import tempfile
import unittest
from contextlib import closing

from openpyxl import load_workbook

import export_excel
import logger


UNIBET = {"key": "unibet_se", "title": "Unibet"}


def make_match(match_id):
    return {"id": match_id, "commence_time": "2026-10-08T18:00:00Z",
            "home_team": f"Hemma {match_id}", "away_team": f"Borta {match_id}"}


class TestExportExcel(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.original_db = logger.DB_FILE
        self.db = os.path.join(self.tmp.name, "test_bets.db")
        self.out = os.path.join(self.tmp.name, "bets_export.xlsx")
        logger.DB_FILE = self.db
        logger.init_db()

    def tearDown(self):
        logger.DB_FILE = self.original_db
        self.tmp.cleanup()

    def add(self, match_id, logged_date, result=None, market_key="h2h", outcome="Hemma",
            point="", bookmaker=UNIBET):
        logger.log_bet(make_match(match_id), "soccer_sweden_allsvenskan", market_key, outcome,
                       point, bookmaker, 2.10, 1.95, 0.0769, 3.5)
        with closing(sqlite3.connect(self.db)) as conn, conn:
            conn.execute("UPDATE bets SET logged_date = ?, result = ? WHERE id = (SELECT MAX(id) FROM bets)",
                         (logged_date, result))

    def export(self):
        count = export_excel.export(self.db, self.out)
        return count, load_workbook(self.out).active

    def test_headers(self):
        _, ws = self.export()
        self.assertEqual([c.value for c in ws[1]], export_excel.HEADERS)
        self.assertEqual(export_excel.HEADERS, [
            "Datum loggad", "Matchdatum", "Sport", "Match", "Lag", "Marknad", "Bookmaker", "Odds",
            "Rättvisa odds", "Edge %", "Units", "Insats Flat (kr)", "Insats Kelly (kr)",
            "Resultat", "V/F Flat", "V/F Kelly"])

    def test_blank_row_between_dates_and_formula_rows(self):
        self.add("m1", "2026-10-06T12:00:00+00:00", "W")
        self.add("m2", "2026-10-06T13:00:00+00:00", "L")
        self.add("m3", "2026-10-07T12:00:00+00:00", "P")
        count, ws = self.export()
        self.assertEqual(count, 3)
        self.assertEqual(ws["D2"].value, "Hemma m1 vs Borta m1")
        self.assertEqual(ws["D3"].value, "Hemma m2 vs Borta m2")
        self.assertTrue(all(c.value is None for c in ws[4]))         # tomrad
        self.assertEqual(ws["D5"].value, "Hemma m3 vs Borta m3")
        self.assertEqual(ws["A2"].value, "2026-10-06")
        self.assertEqual(ws["A5"].value, "2026-10-07")
        self.assertEqual(ws["N5"].value, "P")
        self.assertEqual(ws["O5"].value,
                         '=IF(UPPER(N5)="W",(H5-1)*L5,IF(UPPER(N5)="L",-L5,IF(UPPER(N5)="P",0,"")))')
        self.assertEqual(ws["P5"].value,
                         '=IF(UPPER(N5)="W",(H5-1)*M5,IF(UPPER(N5)="L",-M5,IF(UPPER(N5)="P",0,"")))')

    def test_same_date_no_blank_row(self):
        self.add("m1", "2026-10-06T10:00:00+00:00")
        self.add("m2", "2026-10-06T14:00:00+00:00")
        _, ws = self.export()
        self.assertEqual(ws["D3"].value, "Hemma m2 vs Borta m2")
        self.assertIsNone(ws["N2"].value)

    def test_bookmaker_and_outcome_label(self):
        self.add("m1", "2026-10-06T12:00:00+00:00", market_key="totals", outcome="Over", point=2.5)
        self.add("m2", "2026-10-06T12:00:00+00:00", bookmaker={"key": "okand_bok", "title": "X"})
        _, ws = self.export()
        self.assertEqual(ws["G2"].value, "Unibet (SE)")
        self.assertEqual(ws["E2"].value, "Over 2.5")
        self.assertEqual(ws["G3"].value, "okand_bok")
        self.assertEqual(ws["E3"].value, "Hemma")

    def test_summary_ranges(self):
        self.add("m1", "2026-10-06T12:00:00+00:00", "W")
        self.add("m2", "2026-10-07T12:00:00+00:00", "L")   # rad 4 efter tomrad på rad 3
        _, ws = self.export()
        # Sista datarad 4, tomrad 5, "Summering" rad 6, värden rad 7-12
        self.assertEqual(ws["D6"].value, "Summering")
        self.assertEqual(ws["E7"].value, "=COUNT(H2:H4)")
        self.assertEqual(ws["E8"].value, '=COUNTIF(N2:N4,"W")+COUNTIF(N2:N4,"L")+COUNTIF(N2:N4,"P")')
        self.assertEqual(ws["E9"].value, "=SUM(O2:O4)")
        self.assertEqual(ws["E10"].value, "=SUM(P2:P4)")
        self.assertEqual(ws["E11"].value,
                         '=IFERROR(E9/(SUMIF(N2:N4,"W",L2:L4)+SUMIF(N2:N4,"L",L2:L4)),"")')
        self.assertEqual(ws["E12"].value,
                         '=IFERROR(E10/(SUMIF(N2:N4,"W",M2:M4)+SUMIF(N2:N4,"L",M2:M4)),"")')
        self.assertEqual(ws["E11"].number_format, "0.0%")

    def test_lock_file_refused(self):
        open(os.path.join(self.tmp.name, "~$bets_export.xlsx"), "w").close()
        with self.assertRaisesRegex(export_excel.ExportError, "öppen i Excel"):
            export_excel.export(self.db, self.out)
        self.assertFalse(os.path.exists(self.out))

    def test_protected_file_refused(self):
        target = os.path.join(self.tmp.name, "bets.xlsx")
        with self.assertRaises(export_excel.ExportError):
            export_excel.export(self.db, target)
        self.assertFalse(os.path.exists(target))

    def test_missing_db(self):
        missing = os.path.join(self.tmp.name, "finns_inte.db")
        with self.assertRaisesRegex(export_excel.ExportError, "Hittar inte"):
            export_excel.export(missing, self.out)
        self.assertFalse(os.path.exists(missing))


if __name__ == "__main__":
    unittest.main()
